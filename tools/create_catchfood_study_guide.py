from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\geral\Downloads\PlayLand\Playland")
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)
DOCX_PATH = OUT / "Guia_Defensa_CatchFood.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
PINK = RGBColor(0xE9, 0x6E, 0x9D)
GREEN = RGBColor(0x4F, 0x8D, 0x3A)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT_BLUE = "E8EEF5"
LIGHT_PINK = "FCE8F0"
LIGHT_GREEN = "EAF4E5"
LIGHT_YELLOW = "FFF4CC"
CODE_FILL = "F4F6F9"


def set_font(run, name="Calibri", size=11, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = int(sum(widths_inches) * 1440)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            dxa = int(widths_inches[idx] * 1440)
            cell.width = Inches(widths_inches[idx])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(dxa))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Number"):
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.188)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

header = section.header.paragraphs[0]
header.text = "PLAYLAND · GUÍA DE DEFENSA TÉCNICA · CATCHFOOD"
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(header.runs[0], size=8.5, bold=True, color=MUTED)
add_page_number(section.footer.paragraphs[0])


def add_title(text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_font(r, size=28, bold=True, color=PINK)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(24)
        r2 = p2.add_run(subtitle)
        set_font(r2, size=14, color=DARK_BLUE)


def para(text="", bold_start=None, italic=False, color=None, align=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_start and text.startswith(bold_start):
        r1 = p.add_run(bold_start)
        set_font(r1, bold=True, color=color or DARK)
        r2 = p.add_run(text[len(bold_start):])
        set_font(r2, italic=italic, color=color or DARK)
    else:
        r = p.add_run(text)
        set_font(r, italic=italic, color=color or DARK)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r)
    return p


numbering_counter = 0


def reset_numbering():
    global numbering_counter
    numbering_counter = 0


def numbered(text):
    global numbering_counter
    numbering_counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    r = p.add_run(f"{numbering_counter}.\t{text}")
    set_font(r)
    return p


def callout(label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(label + " ")
    set_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_font(r2, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code(text, caption=None):
    if caption:
        pcap = doc.add_paragraph()
        pcap.paragraph_format.space_before = Pt(4)
        pcap.paragraph_format.space_after = Pt(3)
        rcap = pcap.add_run(caption)
        set_font(rcap, size=9, bold=True, color=MUTED)
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    shade(cell, CODE_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(text.strip("\n").splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, name="Consolas", size=8.3, color=RGBColor(0x18, 0x18, 0x18))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def term_table(rows, title=None):
    if title:
        doc.add_heading(title, level=3)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Término"
    hdr[1].text = "Qué significa y cómo defenderlo"
    for c in hdr:
        shade(c, LIGHT_BLUE)
        for r in c.paragraphs[0].runs:
            set_font(r, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for term, definition in rows:
        cells = table.add_row().cells
        cells[0].text = term
        cells[1].text = definition
        for r in cells[0].paragraphs[0].runs:
            set_font(r, bold=True, color=DARK_BLUE)
        for r in cells[1].paragraphs[0].runs:
            set_font(r)
    set_table_geometry(table, [1.6, 4.9])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def qa(question, short_answer, expanded=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    rq = p.add_run("P: " + question)
    set_font(rq, bold=True, color=BLUE)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    ra = p2.add_run("R: " + short_answer)
    set_font(ra, bold=True, color=DARK)
    if expanded:
        p3 = doc.add_paragraph()
        p3.paragraph_format.left_indent = Inches(0.2)
        p3.paragraph_format.space_after = Pt(5)
        re = p3.add_run("Si piden más detalle: " + expanded)
        set_font(re, size=10, color=MUTED, italic=True)


def normalize_page_breaks_and_numbering(document):
    # Keep only the cover page break. Chapters then flow naturally, avoiding
    # nearly-empty pages caused by a short section plus a forced break.
    page_breaks = document._element.xpath('.//w:br[@w:type="page"]')
    for br in page_breaks[1:]:
        br.getparent().remove(br)



# Cover
add_title("CATCHFOOD", "Guía completa de código para defensa oral")
para("De cero a una explicación técnica segura", bold_start="De cero", align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
callout(
    "Objetivo:",
    "que puedas leer cada construcción importante del proyecto y explicar qué es, qué hace y por qué se utilizó.",
    LIGHT_PINK,
)
para("Incluye Kotlin, Jetpack Compose, ViewModel, Canvas, corrutinas, navegación, Room/SQLite, audio, colisiones y rendimiento.", align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTED, after=8)
para("Preparado a partir del código real de feature/catchfood.", align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTED, italic=True)
doc.add_page_break()

# Study strategy
doc.add_heading("Cómo estudiar esta guía hoy", level=1)
callout("Regla principal:", "no memorices líneas completas. Memorizá responsabilidades y recorridos de datos.", LIGHT_YELLOW)
reset_numbering()
numbered("Leé primero el resumen de 90 segundos y repetilo en voz alta.")
numbered("Aprendé el flujo: UI → ViewModel → GameLogic → UiState → UI → Room.")
numbered("Estudiá el glosario de Kotlin y Compose. Ahí están lambda, callback, state, suspend y demás términos.")
numbered("Revisá los fragmentos línea por línea y explicalos sin mirar.")
numbered("Terminá con las preguntas de jurado y el simulacro oral.")

doc.add_heading("Resumen oral de 90 segundos", level=2)
callout(
    "Respuesta lista para decir:",
    "CatchFood es un minijuego Android hecho con Jetpack Compose. La interfaz está separada del estado mediante CatchFoodViewModel y CatchFoodUiState. El ViewModel coordina el motor CatchFoodGameLogic, que actualiza posiciones, velocidad, colisiones y reglas de derrota. El loop se sincroniza con los frames reales usando withFrameNanos y delta time. Para evitar lag, los objetos se dibujan juntos en un Canvas, en lugar de usar muchos componentes Image que se reubiquen cada frame. Al terminar una partida, el ViewModel guarda el puntaje mediante un Repository en Room, que utiliza SQLite local. El histórico ejecuta una consulta SQL ordenada y muestra los diez mejores resultados.",
    LIGHT_GREEN,
)

doc.add_heading("Mapa mental", level=2)
para("Entrada táctil → lambda de arrastre → ViewModel.movePlayer() → GameLogic.playerX → Canvas redibuja el conejo.")
para("Frame nuevo → ViewModel.updateGame() → GameLogic.updateGame() → colisión/puntaje → UiState cambia → Compose actualiza contadores.")
para("Derrota → ViewModel captura finalScore → corrutina → Repository → DAO → Room → SQLite.")
para("Histórico → produceState → Repository → DAO → SELECT ORDER BY LIMIT 10 → LazyColumn.")
doc.add_page_break()

# Fundamentals
doc.add_heading("1. Fundamentos de Kotlin que aparecen en CatchFood", level=1)
term_table([
    ("val", "Declara una referencia que no puede reasignarse. El objeto interno sí podría ser mutable."),
    ("var", "Declara una variable cuyo valor puede cambiar."),
    ("fun", "Declara una función: una unidad de código que puede recibir parámetros y devolver un resultado."),
    ("class", "Define un tipo con datos y comportamiento."),
    ("data class", "Clase orientada a datos. Kotlin genera equals, hashCode, toString, componentN y copy."),
    ("object", "Crea una única instancia. En navegación funciona como destino singleton."),
    ("private", "El elemento solo puede usarse dentro del alcance permitido, por ejemplo la misma clase o archivo."),
    ("override", "Indica que una propiedad o función reemplaza una declaración heredada."),
    ("null", "Representa ausencia de valor."),
    ("?", "Convierte un tipo en nullable: puede contener un valor o null."),
    ("?:", "Operador Elvis: usa el valor derecho si el izquierdo es null."),
    ("?.", "Llamada segura: ejecuta el acceso solo si el valor no es null."),
    ("::", "Referencia a función o propiedad; no la ejecuta todavía."),
    ("by", "Delegación. Permite leer/escribir un State como si fuera una variable normal."),
], "Palabras reservadas y operadores")

doc.add_heading("Lambda", level=2)
para("Una lambda es una función sin nombre que puede guardarse, pasarse como parámetro o ejecutarse después.")
code("""
onClick = { onGoToGame() }
""", "Ejemplo real del menú")
bullet("`onClick` espera una función que no recibe parámetros y no devuelve información útil: `() -> Unit`.")
bullet("Las llaves `{ ... }` contienen el cuerpo de la lambda.")
bullet("La lambda no se ejecuta al construir la pantalla; se ejecuta cuando el usuario toca el botón.")
callout("Respuesta oral:", "Una lambda es una función anónima que puedo pasar como dato. En CatchFood la uso para comunicar eventos como clics, volver o mover al jugador.", LIGHT_GREEN)

doc.add_heading("Callback", level=2)
para("Un callback es una función que se entrega a otro componente para que la llame cuando ocurra un evento.")
code("""
fun CatchFoodScreen(
    onBack: () -> Unit,
    onGoToGame: () -> Unit,
    onGoToLeaderboard: () -> Unit
)
""")
bullet("La pantalla no conoce NavController.")
bullet("Solo sabe que puede llamar `onGoToGame()`.")
bullet("El NavHost decide qué navegación concreta ocurre.")
callout("Ventaja:", "reduce el acoplamiento y hace que la pantalla sea más reutilizable y fácil de probar.")

doc.add_heading("Unit", level=2)
para("`Unit` es el equivalente conceptual a `void`: indica que la función termina una acción pero no devuelve un resultado relevante.")

doc.add_heading("Genéricos", level=2)
para("Los símbolos `<T>` permiten que una clase o función trabaje con distintos tipos. Ejemplo: `List<Int>` es una lista cuyo elemento debe ser Int; `List<FallingObject>` contiene objetos que caen.")

doc.add_heading("Funciones de orden superior", level=2)
para("Son funciones que reciben o devuelven funciones. `forEach`, `use`, `remember`, `LaunchedEffect` y `detectDragGestures` reciben lambdas.")
doc.add_page_break()

# Compose
doc.add_heading("2. Jetpack Compose desde cero", level=1)
term_table([
    ("@Composable", "Anotación que marca una función capaz de describir interfaz declarativa."),
    ("Composición", "Proceso inicial en el que Compose ejecuta funciones y construye el árbol de UI."),
    ("Recomposición", "Reejecución selectiva cuando cambia un estado observado."),
    ("State", "Valor observable. Cuando cambia, Compose invalida la fase donde fue leído."),
    ("Modifier", "Cadena de instrucciones para tamaño, posición, entrada táctil, fondo y otros comportamientos."),
    ("dp", "Unidad independiente de densidad usada para tamaños y posiciones visuales."),
    ("sp", "Unidad para texto que considera densidad y preferencias de accesibilidad."),
    ("Context", "Acceso a recursos y servicios Android; aquí se usa para assets, MediaPlayer y Room."),
])

doc.add_heading("Qué significa UI declarativa", level=2)
para("En una UI imperativa ordenarías manualmente: busca el Text y cambia su contenido. En Compose declarás cómo debe verse la UI para un estado dado:")
code('Text("Comida: ${uiState.score}")')
para("Cuando `uiState.score` cambia, Compose vuelve a ejecutar la parte necesaria y genera el texto nuevo.")

doc.add_heading("remember", level=2)
code("""
val bitmaps = remember(context, density) {
    CatchFoodBitmaps.load(context, density)
}
""")
bullet("`remember` conserva el resultado entre recomposiciones.")
bullet("`context` y `density` son claves. Si cambia una, se vuelve a calcular.")
bullet("Sin remember, los PNG podrían decodificarse repetidamente.")
callout("Respuesta oral:", "remember evita repetir trabajo durante recomposiciones, pero no reemplaza al ViewModel ni garantiza persistencia después de destruir la pantalla.")

doc.add_heading("Estado y delegación by", level=2)
code("""
var uiState by mutableStateOf(gameLogic.toUiState())
    private set
""")
bullet("`mutableStateOf` crea un contenedor observable.")
bullet("`by` delega get y set: permite escribir `uiState.score` en vez de `uiState.value.score`.")
bullet("`private set` permite lectura pública, pero solo el ViewModel puede cambiarlo.")

doc.add_heading("Efectos", level=2)
term_table([
    ("LaunchedEffect", "Inicia una corrutina ligada a la vida del Composable. Se cancela cuando sale de composición o cambia su clave."),
    ("DisposableEffect", "Ejecuta inicialización y exige limpieza mediante onDispose."),
    ("produceState", "Ejecuta trabajo suspendido y expone el resultado como State de Compose."),
])
doc.add_page_break()

# Architecture/files
doc.add_heading("3. Arquitectura real del módulo", level=1)
para("El módulo no implementa MVVM de forma académicamente pura, pero aplica una separación clara inspirada en MVVM y Repository:")
bullet("View: CatchFoodScreen, CatchFoodGame y CatchFoodLeaderboard.")
bullet("ViewModel: CatchFoodViewModel.")
bullet("Model/dominio: CatchFoodGameLogic, FallingObject y CatchFoodUiState.")
bullet("Datos: Entity, DAO, Database y Repository.")
bullet("Servicios: CatchFoodAudioManager.")
bullet("Coordinación: CatchFoodNavHost y destinos.")

callout("Frase de defensa:", "La UI no guarda directamente en SQLite ni implementa las reglas. Envía eventos al ViewModel, observa UiState y delega persistencia al Repository.", LIGHT_GREEN)

doc.add_heading("Flujo de dependencias", level=2)
reset_numbering()
numbered("CatchFoodGame obtiene el ViewModel con `viewModel()`.")
numbered("El ViewModel posee una instancia de CatchFoodGameLogic.")
numbered("GameLogic actualiza el mundo del juego.")
numbered("El ViewModel sincroniza valores visibles hacia CatchFoodUiState.")
numbered("Compose lee UiState y dibuja contadores/menús.")
numbered("Al perder, el ViewModel usa Repository para guardar en Room.")

doc.add_heading("Por qué separar responsabilidades", level=2)
bullet("Cambiar el diseño no obliga a reescribir SQL.")
bullet("Cambiar la base no obliga a reescribir el Canvas.")
bullet("Las reglas pueden estudiarse o probarse de forma aislada.")
bullet("El estado no desaparece por una simple recomposición.")
doc.add_page_break()

# Menu
doc.add_heading("4. Menú principal línea por línea", level=1)
code("""
@Composable
fun CatchFoodScreen(
    onBack: () -> Unit,
    onGoToGame: () -> Unit,
    onGoToLeaderboard: () -> Unit
) {
""")
term_table([
    ("@Composable", "Le dice al compilador de Compose que la función puede emitir interfaz."),
    ("fun", "Inicia la declaración de una función."),
    ("CatchFoodScreen", "Nombre de la función/pantalla."),
    ("onBack", "Parámetro que almacena una función para volver."),
    (":", "Separa el nombre del parámetro de su tipo."),
    ("() -> Unit", "Tipo función: cero parámetros, devuelve Unit."),
])

doc.add_heading("Contexto y audio", level=2)
code("""
val context = LocalContext.current
val audioManager = remember { CatchFoodAudioManager(context) }

DisposableEffect(Unit) {
    audioManager.startMusic()
    onDispose { audioManager.stopMusic() }
}
""")
bullet("`LocalContext.current` obtiene el Context disponible en el árbol Compose.")
bullet("`remember` mantiene el mismo AudioManager durante recomposiciones.")
bullet("`DisposableEffect(Unit)` tiene una clave constante: se inicia una vez para esa entrada en composición.")
bullet("`onDispose` se ejecuta al abandonar el menú.")

doc.add_heading("Carga segura del fondo", level=2)
code("""
val fondoCompleto = remember {
    try {
        context.assets.open(
            "games/catchfood/sprites/background_historico.png"
        ).use { inputStream ->
            BitmapFactory.decodeStream(inputStream).asImageBitmap()
        }
    } catch (e: Exception) {
        null
    }
}
""")
bullet("`try/catch` evita cerrar la aplicación si falla la lectura.")
bullet("`assets.open` abre un InputStream.")
bullet("`use` garantiza que el stream se cierre incluso si ocurre una excepción.")
bullet("`decodeStream` crea un Bitmap de Android.")
bullet("`asImageBitmap` lo adapta para Compose.")
bullet("El tipo inferido es `ImageBitmap?` porque el catch devuelve null.")

doc.add_heading("let y null safety", level=2)
code("""
fondoCompleto?.let { bitmap ->
    Image(bitmap = bitmap, ...)
}
""")
para("`?.let` ejecuta la lambda solo si fondoCompleto no es null. Dentro de la lambda, `bitmap` ya no es nullable.")
doc.add_page_break()

# Navigation
doc.add_heading("5. Navegación y lambdas", level=1)
code("""
@Serializable
object CatchFoodGameScreen
""")
bullet("`@Serializable` permite convertir el destino a una representación navegable.")
bullet("`object` crea una única instancia porque la ruta no necesita parámetros.")

code("""
composable<CatchFoodGameScreen> {
    CatchFoodGame(
        onBack = { navController.popBackStack() }
    )
}
""")
bullet("`composable<T>` registra una pantalla para el tipo T.")
bullet("La lambda exterior es el contenido que Navigation mostrará.")
bullet("La lambda `onBack` se pasa como callback.")
bullet("`popBackStack()` elimina la pantalla actual y revela la anterior.")

callout("Pregunta trampa:", "¿Por qué no pasar NavController a todas las pantallas? Porque aumentaría el acoplamiento. Las pantallas solo necesitan expresar eventos; el NavHost decide las rutas.", LIGHT_YELLOW)
doc.add_page_break()

# UiState
doc.add_heading("6. CatchFoodUiState palabra por palabra", level=1)
code("""
data class CatchFoodUiState(
    val score: Int = 0,
    val poisonHits: Int = 0,
    val missedFood: Int = 0,
    val playerState: String = "eat",
    val isPaused: Boolean = false,
    val isGameOver: Boolean = false
)
""")
term_table([
    ("data class", "Clase de datos. Kotlin genera operaciones útiles, especialmente copy()."),
    ("val", "Cada propiedad es inmutable dentro de una instancia concreta."),
    ("Int", "Número entero de 32 bits."),
    ("String", "Cadena de caracteres; aquí identifica el sprite del conejo."),
    ("Boolean", "Solo puede ser true o false."),
    ("= 0 / false", "Valor por defecto usado cuando se construye sin argumento."),
])

doc.add_heading("Qué es UiState", level=2)
para("UiState significa User Interface State: el estado necesario para representar la pantalla en un instante.")
bullet("No es la base de datos.")
bullet("No es la pantalla.")
bullet("No contiene funciones de negocio.")
bullet("Es una fotografía inmutable de valores visibles.")

code("""
uiState = uiState.copy(isPaused = true)
""")
para("`copy` crea otro CatchFoodUiState copiando todo, excepto isPaused. Esta inmutabilidad hace predecibles los cambios de UI.")
callout("Respuesta oral:", "UiState es un único objeto que agrupa todo lo que la pantalla necesita observar. Cuando creo una copia con valores nuevos, Compose detecta el cambio y recompone lo necesario.", LIGHT_GREEN)
doc.add_page_break()

# ViewModel
doc.add_heading("7. CatchFoodViewModel a nivel de código", level=1)
code("""
class CatchFoodViewModel(application: Application)
    : AndroidViewModel(application) {
""")
bullet("La clase recibe `Application`, un contexto con vida igual al proceso de la app.")
bullet("Los dos puntos indican herencia.")
bullet("Hereda de AndroidViewModel porque necesita contexto para Room.")
bullet("No guarda una Activity, evitando fugas de memoria.")

code("""
val gameLogic = CatchFoodGameLogic()
private val scoreRepository = CatchFoodScoreRepository(application)
private var scoreSavedForCurrentGame = false
""")
bullet("El ViewModel posee el motor del juego.")
bullet("Repository es privado: la UI no accede directamente a datos.")
bullet("El booleano funciona como guardia contra inserciones duplicadas.")

doc.add_heading("Actualización y guardado", level=2)
code("""
fun updateGame(deltaSeconds: Float) {
    gameLogic.updateGame(deltaSeconds)
    syncUiState()

    if (gameLogic.isGameOver && !scoreSavedForCurrentGame) {
        scoreSavedForCurrentGame = true
        val finalScore = gameLogic.score
        viewModelScope.launch {
            scoreRepository.saveScore(finalScore)
        }
    }
}
""")
bullet("`deltaSeconds: Float` es tiempo transcurrido con decimales.")
bullet("Primero avanza el motor; después sincroniza datos visibles.")
bullet("`&&` significa AND: ambas condiciones deben cumplirse.")
bullet("`!` niega un Boolean. `!scoreSaved...` significa que todavía no fue guardado.")
bullet("`finalScore` captura el valor antes de iniciar la corrutina; un reinicio rápido no lo cambia.")
bullet("`viewModelScope.launch` inicia una corrutina ligada al ViewModel.")

doc.add_heading("Qué es una corrutina", level=2)
para("Es una unidad de trabajo suspendible y liviana. Permite esperar operaciones como SQLite sin bloquear la interfaz.")
callout("No digas:", "‘una corrutina es un hilo’. Puede usar hilos, pero es una abstracción más liviana y puede suspenderse sin bloquearlos.", LIGHT_YELLOW)

doc.add_heading("Sincronización eficiente", level=2)
para("`syncUiState()` compara cada campo antes de crear un estado nuevo. Así no genera objetos ni recomposiciones en cada frame cuando solo cambian posiciones del Canvas.")
code("""
if (uiState.score != gameLogic.score || ...) {
    uiState = gameLogic.toUiState(isPaused = uiState.isPaused)
}
""")
bullet("`!=` significa diferente.")
bullet("`||` significa OR: basta que una condición sea verdadera.")
bullet("`toUiState` es una función de extensión privada sobre GameLogic.")
doc.add_page_break()

# Engine
doc.add_heading("8. Modelo y motor del juego", level=1)
code("""
data class FallingObject(
    var x: Float,
    var y: Float,
    val isPoison: Boolean
)
""")
bullet("x/y son var porque cambian mientras cae o reaparece.")
bullet("isPoison es val porque el tipo del objeto no cambia durante su vida.")

doc.add_heading("Estado optimizado de Compose", level=2)
code("""
var playerX by mutableFloatStateOf(300f)
var score by mutableIntStateOf(0)
var isGameOver by mutableStateOf(false)
""")
bullet("Las versiones especializadas Float/Int evitan boxing innecesario.")
bullet("`300f` es un literal Float; la `f` evita que Kotlin lo interprete como Double.")
bullet("Compose observa estos valores en la fase donde se leen.")

doc.add_heading("private set", level=2)
code("""
var speed = INITIAL_SPEED
    private set
""")
para("Cualquier código puede leer speed, pero solo CatchFoodGameLogic puede cambiarla.")

doc.add_heading("Companion object y constantes", level=2)
code("""
private companion object {
    const val TARGET_FPS = 60f
    const val INITIAL_SPEED = 8f
}
""")
bullet("`companion object` contiene miembros asociados a la clase, no a una instancia específica.")
bullet("`const val` es una constante conocida en compilación.")
bullet("Los nombres en mayúsculas siguen la convención de constantes.")
doc.add_page_break()

# Loop
doc.add_heading("9. Loop, FPS y delta time", level=1)
code("""
LaunchedEffect(gameViewModel) {
    var previousFrameNanos = 0L
    while (true) {
        withFrameNanos { frameNanos ->
            // actualización
        }
    }
}
""")
bullet("LaunchedEffect inicia el loop como corrutina y lo cancela al salir de la pantalla.")
bullet("`0L` es un literal Long; la L identifica 64 bits.")
bullet("`while (true)` crea el ciclo continuo.")
bullet("`withFrameNanos` suspende hasta que Android vaya a producir otro frame.")

code("""
val deltaSeconds =
    (frameNanos - previousFrameNanos) / 1_000_000_000f
""")
bullet("Los timestamps vienen en nanosegundos.")
bullet("Un segundo contiene mil millones de nanosegundos.")
bullet("Los guiones bajos solo mejoran lectura del número.")

code("""
val frameScale =
    (deltaSeconds * TARGET_FPS).coerceIn(0f, MAX_FRAME_SCALE)

obj.y += speed * frameScale
""")
para("A 60 FPS, delta ≈ 0.0167 s y frameScale ≈ 1. En un frame más largo, el desplazamiento aumenta proporcionalmente.")
callout("Respuesta oral:", "Delta time hace que el movimiento dependa del tiempo real y no del número de frames. coerceIn limita saltos si el dispositivo se congela momentáneamente.", LIGHT_GREEN)

doc.add_heading("Pausa", level=2)
para("Cuando UiState indica pausa o derrota, el loop no llama updateGame y reinicia previousFrameNanos. Al continuar no acumula todo el tiempo pausado como un salto enorme.")
doc.add_page_break()

# Spawn
doc.add_heading("10. Spawn y cola de objetos", level=1)
code("""
val objectTypes =
    (List(8) { false } + List(2) { true }).shuffled()
""")
bullet("`List(8) { false }` fabrica ocho valores false: zanahorias.")
bullet("`+` concatena ambas listas.")
bullet("`List(2) { true }` representa dos venenos.")
bullet("`shuffled()` devuelve una copia con orden aleatorio.")

code("""
objectTypes.forEachIndexed { index, isPoison ->
    objects.add(
        FallingObject(
            x = Random.nextInt(80, 850).toFloat(),
            y = INITIAL_SPAWN_Y - SPAWN_SPACING_WORLD * index,
            isPoison = isPoison
        )
    )
}
""")
bullet("`forEachIndexed` entrega valor e índice.")
bullet("Los parámetros de la lambda son index e isPoison.")
bullet("Random.nextInt usa mínimo incluido y máximo excluido.")
bullet("toFloat convierte Int a Float.")
bullet("Cada índice resta otra separación vertical.")

doc.add_heading("Reaparición detrás de la cola", level=2)
code("""
val objectBehindQueue = objects
    .asSequence()
    .filter { it !== obj }
    .minOfOrNull { it.y }
    ?: INITIAL_SPAWN_Y
""")
bullet("`asSequence` procesa operaciones de forma perezosa.")
bullet("`it` es el nombre implícito del único parámetro de una lambda.")
bullet("`!==` compara identidad: excluye exactamente la misma instancia.")
bullet("`minOfOrNull` obtiene la Y más pequeña o null si no hay elementos.")
bullet("`?:` usa INITIAL_SPAWN_Y si el resultado fue null.")
para("Luego se resta SPAWN_SPACING_WORLD. Por eso el objeto reaparece después del último y no encima de otro.")
doc.add_page_break()

# Collision
doc.add_heading("11. Colisiones explicadas", level=1)
para("El juego usa AABB: Axis-Aligned Bounding Box. Son rectángulos sin rotación.")
code("""
val touchesPlayer = viewportHeightDp > 0f &&
    objectLeft < playerLeft + PLAYER_HITBOX_WIDTH_DP &&
    objectLeft + OBJECT_HITBOX_WIDTH_DP > playerLeft &&
    objectTop + OBJECT_HITBOX_INSET_TOP_DP <
        playerTop + PLAYER_HITBOX_HEIGHT_DP &&
    objectTop + OBJECT_HITBOX_INSET_TOP_DP +
        OBJECT_HITBOX_HEIGHT_DP > playerTop
""")
term_table([
    ("Condición 1", "El lado izquierdo del objeto está antes del lado derecho del jugador."),
    ("Condición 2", "El lado derecho del objeto pasó el lado izquierdo del jugador."),
    ("Condición 3", "La parte superior del objeto está antes de la parte inferior del jugador."),
    ("Condición 4", "La parte inferior del objeto pasó la parte superior del jugador."),
])
para("Si las cuatro condiciones son verdaderas, existe superposición horizontal y vertical.")

doc.add_heading("Por qué hay insets", level=2)
para("Los PNG ocupan una caja grande y contienen transparencia. Los insets mueven y reducen la hitbox hacia la parte visible. Esto evita recoger veneno cuando el dibujo parece pasar al lado.")

code("""
if (touchesPlayer) {
    if (obj.isPoison) {
        poisonHits++
    } else {
        score++
    }
    resetObject(obj)
    return@forEach
}
""")
bullet("`++` incrementa en uno.")
bullet("`return@forEach` termina la iteración actual, no updateGame completa.")
bullet("Esto evita que el mismo objeto sea contado también como perdido en el mismo frame.")

doc.add_heading("Condiciones de derrota", level=2)
bullet("Tres venenos: playerState = dead e isGameOver = true.")
bullet("Cinco zanahorias perdidas: playerState = dead e isGameOver = true.")
doc.add_page_break()

# Canvas
doc.add_heading("12. Canvas y rendimiento", level=1)
code("""
Canvas(
    modifier = Modifier
        .fillMaxSize()
        .pointerInput(gameLogic) { ... }
) {
    gameLogic.frameVersion
    // drawImage(...)
}
""")
bullet("Canvas proporciona un DrawScope para dibujar directamente.")
bullet("fillMaxSize ocupa la pantalla disponible.")
bullet("pointerInput instala manejo de gestos.")
bullet("gameLogic es la clave; si cambia, se reinicia ese manejador.")

doc.add_heading("frameVersion", level=2)
para("Los x/y internos de FallingObject son variables normales por rendimiento. Al final de cada update, frameVersion aumenta. Como se lee dentro del bloque de dibujo, Compose invalida el Canvas y solicita redraw sin reconstruir toda la pantalla.")

code("""
drawImage(
    image = it,
    dstOffset = IntOffset(x, y),
    dstSize = IntSize(objectSize, objectSize)
)
""")
term_table([
    ("image", "Bitmap que se dibuja."),
    ("dstOffset", "Coordenada entera de destino, esquina superior izquierda."),
    ("dstSize", "Ancho y alto en píxeles del rectángulo de destino."),
    ("density", "Cantidad aproximada de píxeles físicos por dp."),
])

callout("Por qué no diez Image:", "Mover diez composables Image cada frame provoca más recomposición, medición y colocación. Un Canvas agrupa el trabajo en una sola fase de dibujo.", LIGHT_GREEN)

doc.add_heading("Gestos", level=2)
code("""
detectDragGestures { _, dragAmount ->
    onMovePlayer(dragAmount.x)
}
""")
bullet("El guion bajo ignora un parámetro que no se necesita.")
bullet("dragAmount es un desplazamiento desde el evento anterior.")
bullet("Solo se usa x porque el conejo se mueve horizontalmente.")
doc.add_page_break()

# Images
doc.add_heading("13. Bitmaps y memoria", level=1)
para("CatchFoodBitmaps es un data class privado que agrupa referencias a sprites ya decodificados.")
code("""
private data class CatchFoodBitmaps(
    val player: ImageBitmap?,
    val food: ImageBitmap?,
    val poison: ImageBitmap?,
    ...
)
""")
bullet("Es private porque solo CatchFoodGame.kt necesita conocerlo.")
bullet("Cada propiedad es nullable para tolerar errores de assets.")

doc.add_heading("Decodificación en dos pasos", level=2)
code("""
val bounds = BitmapFactory.Options().apply {
    inJustDecodeBounds = true
}
""")
bullet("apply ejecuta una lambda usando el objeto como receptor.")
bullet("inJustDecodeBounds lee dimensiones sin cargar todos los píxeles.")

code("""
val options = BitmapFactory.Options().apply {
    inSampleSize = sampleSize
    inPreferredConfig = Bitmap.Config.ARGB_8888
}
""")
bullet("inSampleSize reduce ancho y alto durante decodificación.")
bullet("ARGB_8888 usa 8 bits para alfa, rojo, verde y azul.")
bullet("Reducir resolución evita ocupar memoria con píxeles invisibles.")

doc.add_heading("runCatching", level=2)
para("`runCatching { ... }.getOrNull()` convierte cualquier excepción del bloque en null. La UI puede omitir el sprite en vez de cerrarse.")
doc.add_page_break()

# Room
doc.add_heading("14. Room y SQLite desde cero", level=1)
callout("Respuesta esencial:", "Room no es PostgreSQL. Room genera una capa segura sobre SQLite, una base local almacenada dentro del teléfono.", LIGHT_YELLOW)

doc.add_heading("Entity = tabla", level=2)
code("""
@Entity(tableName = "catch_food_scores")
data class CatchFoodScoreEntity(
    @PrimaryKey(autoGenerate = true)
    val id: Long = 0,
    val foodCaught: Int
)
""")
term_table([
    ("@Entity", "Anotación Room que convierte la clase en tabla."),
    ("tableName", "Nombre real de la tabla SQLite."),
    ("@PrimaryKey", "Columna que identifica de forma única una fila."),
    ("autoGenerate", "SQLite asigna automáticamente el siguiente id."),
    ("Long", "Entero de 64 bits, apropiado para identificadores."),
])

doc.add_heading("DAO = operaciones", level=2)
code("""
@Dao
interface CatchFoodScoreDao {
    @Insert
    suspend fun insert(score: CatchFoodScoreEntity)
}
""")
bullet("interface define un contrato sin implementación manual.")
bullet("Room/KSP genera la implementación.")
bullet("@Insert indica el SQL de inserción automáticamente.")
bullet("suspend permite suspender la corrutina durante I/O.")

doc.add_heading("Consulta SQL", level=2)
code("""
SELECT foodCaught
FROM catch_food_scores
ORDER BY foodCaught DESC, id ASC
LIMIT 10
""")
term_table([
    ("SELECT", "Columnas que se desean devolver."),
    ("FROM", "Tabla que se consulta."),
    ("ORDER BY", "Orden de resultados."),
    ("DESC", "De mayor a menor."),
    ("ASC", "De menor a mayor; desempata por partida más antigua."),
    ("LIMIT 10", "Devuelve como máximo diez filas."),
])
doc.add_page_break()

# Database/repository
doc.add_heading("15. Database, Singleton y Repository", level=1)
code("""
@Database(
    entities = [CatchFoodScoreEntity::class],
    version = 1,
    exportSchema = false
)
abstract class CatchFoodDatabase : RoomDatabase()
""")
bullet("`entities = [...]` registra las tablas.")
bullet("`::class` obtiene la referencia Kotlin a la clase.")
bullet("version identifica la versión del esquema.")
bullet("Si el esquema cambia, debe incrementarse y normalmente añadirse una migración.")
bullet("abstract permite que Room genere la implementación concreta.")

doc.add_heading("Singleton thread-safe", level=2)
code("""
@Volatile
private var instance: CatchFoodDatabase? = null

fun getInstance(context: Context): CatchFoodDatabase =
    instance ?: synchronized(this) {
        instance ?: Room.databaseBuilder(...).build()
            .also { instance = it }
    }
""")
term_table([
    ("@Volatile", "Garantiza que los hilos vean el valor actualizado de instance."),
    ("synchronized", "Solo un hilo puede ejecutar el bloque simultáneamente."),
    ("?:", "Si instance existe la devuelve; si no, crea la base."),
    ("also", "Devuelve el mismo objeto y ejecuta una acción lateral: guardarlo en instance."),
    ("applicationContext", "Contexto de larga vida; evita conservar una Activity."),
])

doc.add_heading("Repository", level=2)
code("""
class CatchFoodScoreRepository(context: Context) {
    private val scoreDao =
        CatchFoodDatabase.getInstance(context).scoreDao()

    suspend fun saveScore(score: Int) {
        scoreDao.insert(CatchFoodScoreEntity(foodCaught = score))
    }
}
""")
para("El Repository ofrece operaciones con nombres del dominio y oculta DAO/Room al ViewModel. Esto permite cambiar la fuente de datos con menor impacto.")
doc.add_page_break()

# Leaderboard
doc.add_heading("16. Histórico y estado asíncrono", level=1)
code("""
val scores by produceState(
    initialValue = emptyList(),
    scoreRepository
) {
    value = scoreRepository.getTopScores()
}
""")
bullet("produceState devuelve State<List<Int>>.")
bullet("by permite usar scores directamente.")
bullet("emptyList es el valor mientras termina SQLite.")
bullet("scoreRepository es una clave del efecto.")
bullet("value actualiza el State y provoca recomposición.")

doc.add_heading("LazyColumn", level=2)
code("""
LazyColumn {
    itemsIndexed(scores) { index, score ->
        Text("${index + 1}.")
        Text(score.toString())
    }
}
""")
bullet("LazyColumn es una lista vertical que compone elementos según se necesitan.")
bullet("itemsIndexed entrega posición y valor.")
bullet("Interpolación `${...}` inserta expresiones en String.")
bullet("toString convierte Int a texto.")

callout("Flujo completo:", "Room devuelve List<Int> → produceState cambia scores → Compose recompone → LazyColumn dibuja ranking.", LIGHT_GREEN)
doc.add_page_break()

# Audio/lifecycle
doc.add_heading("17. Audio y ciclo de vida", level=1)
code("""
mediaPlayer = MediaPlayer().apply {
    context.assets.openFd("...mp3").use { fd ->
        setDataSource(fd.fileDescriptor, fd.startOffset, fd.length)
    }
    isLooping = true
    prepare()
    start()
}
""")
bullet("MediaPlayer reproduce audio nativo Android.")
bullet("openFd obtiene descriptor, desplazamiento y longitud del asset.")
bullet("setDataSource define el origen.")
bullet("isLooping repite la canción.")
bullet("prepare prepara sincrónicamente.")
bullet("start inicia reproducción.")

code("""
mediaPlayer?.apply {
    if (isPlaying) stop()
    release()
}
mediaPlayer = null
""")
bullet("`?.` evita operar si mediaPlayer es null.")
bullet("release libera recursos nativos.")
bullet("Asignar null permite crear uno nuevo al volver.")
callout("Defensa:", "La limpieza está conectada a onDispose para evitar audio duplicado y fugas de recursos.", LIGHT_GREEN)
doc.add_page_break()

# Full flow
doc.add_heading("18. Recorrido completo de una partida", level=1)
reset_numbering()
for item in [
    "El NavHost muestra CatchFoodMenu como destino inicial.",
    "CatchFoodScreen carga el fondo e inicia la música.",
    "El usuario pulsa Jugar; la lambda llama navController.navigate.",
    "CatchFoodGame obtiene CatchFoodViewModel mediante viewModel().",
    "El ViewModel ya posee GameLogic y UiState inicial.",
    "LaunchedEffect inicia el loop de frames.",
    "withFrameNanos entrega el timestamp del nuevo frame.",
    "Se calcula deltaSeconds y ViewModel llama updateGame.",
    "GameLogic mueve objetos y aumenta gradualmente speed.",
    "GameLogic calcula hitboxes y detecta colisiones.",
    "El ViewModel sincroniza cambios visibles hacia UiState.",
    "Compose actualiza contadores o la imagen del conejo.",
    "frameVersion invalida el Canvas para nuevas posiciones.",
    "El usuario arrastra; la lambda envía dragAmount.x al ViewModel.",
    "Con tres venenos o cinco comidas perdidas se activa game over.",
    "El ViewModel captura el puntaje y lanza una corrutina.",
    "Repository crea Entity y DAO inserta en SQLite.",
    "Al abrir Histórico, DAO ejecuta SELECT ORDER BY LIMIT 10.",
    "produceState recibe la lista y LazyColumn muestra los resultados.",
]:
    numbered(item)

doc.add_heading("La explicación de una sola zanahoria", level=2)
para("Se crea con x aleatoria e y negativa. En cada frame y aumenta. Canvas convierte sus coordenadas y la dibuja. Cuando su hitbox se superpone con la del conejo, score aumenta, UiState se sincroniza, el Text se recompone y la zanahoria se coloca detrás de la cola.")
doc.add_page_break()

# Glossary
doc.add_heading("19. Glosario rápido para preguntas sorpresa", level=1)
term_table([
    ("Anotación", "Metadato con @ que informa a compiladores o librerías, por ejemplo @Composable, @Entity y @Dao."),
    ("API", "Contrato que permite que componentes se comuniquen; no significa necesariamente Internet."),
    ("Callback", "Función entregada para ejecutarse cuando ocurra algo."),
    ("Lambda", "Función anónima expresada entre llaves."),
    ("Scope", "Alcance donde existe un nombre o donde vive una corrutina."),
    ("Corrutina", "Trabajo suspendible y liviano para operaciones concurrentes."),
    ("Suspend", "Función que puede pausar su ejecución sin bloquear el hilo."),
    ("Hilo principal", "Hilo que procesa UI; bloquearlo provoca congelamiento o ANR."),
    ("Estado", "Información mutable que determina comportamiento o apariencia."),
    ("UiState", "Fotografía inmutable de los datos visibles de una pantalla."),
    ("Recomposición", "Reejecución selectiva de Composables por cambios observados."),
    ("Inmutabilidad", "Crear una nueva instancia en vez de alterar la anterior."),
    ("Acoplamiento", "Grado de dependencia entre componentes."),
    ("Encapsulación", "Ocultar detalles y exponer solo operaciones necesarias."),
    ("Singleton", "Patrón que mantiene una única instancia compartida."),
    ("Repository", "Abstracción que oculta las fuentes y operaciones de datos."),
    ("DAO", "Interfaz de acceso a datos con operaciones de lectura/escritura."),
    ("Entity", "Clase que Room transforma en tabla."),
    ("Persistencia", "Datos que sobreviven al cierre de la aplicación."),
    ("SQLite", "Base relacional embebida dentro del dispositivo."),
    ("KSP", "Procesador de símbolos que genera código de Room al compilar."),
    ("Canvas", "Superficie de dibujo directo para gráficos eficientes."),
    ("FPS", "Frames dibujados por segundo."),
    ("Delta time", "Tiempo transcurrido desde la actualización anterior."),
    ("Hitbox", "Área matemática usada para colisiones."),
    ("AABB", "Colisión entre rectángulos alineados con los ejes."),
    ("Nullable", "Tipo que admite null."),
    ("Leak", "Recurso retenido cuando ya debería liberarse."),
], "Conceptos en una frase")
doc.add_page_break()

# Q&A
doc.add_heading("20. Banco de preguntas de defensa", level=1)
qa("¿Qué patrón de arquitectura utilizaste?", "Una separación inspirada en MVVM con Repository.", "Compose actúa como View, CatchFoodViewModel administra UiState y eventos, GameLogic contiene reglas y Repository abstrae Room.")
qa("¿Qué es una lambda?", "Una función anónima que se puede pasar como dato.", "Los botones reciben lambdas onClick y las pantallas reciben callbacks de navegación.")
qa("¿Qué es UiState?", "Un objeto inmutable con todo lo que la pantalla necesita mostrar.", "CatchFoodUiState contiene puntajes, estado del jugador, pausa y derrota.")
qa("¿Por qué ViewModel?", "Para separar lógica/estado de la UI y conservarlos ante recomposiciones.", "También ofrece viewModelScope para guardar en Room sin bloquear la interfaz.")
qa("¿Qué es recomposición?", "La reejecución selectiva de UI cuando cambia un State observado.")
qa("¿Por qué Canvas?", "Para redibujar objetos animados eficientemente en una sola superficie.", "Evita medir y colocar diez Image composables en cada frame.")
qa("¿Por qué delta time?", "Para que la velocidad dependa del tiempo real y no de los FPS.")
qa("¿Cómo detectás colisión?", "Con cuatro comparaciones AABB entre hitboxes rectangulares reducidas.")
qa("¿Por qué la hitbox es menor que el PNG?", "Porque el PNG contiene transparencia y una caja visual mayor al personaje real.")
qa("¿Cómo evitás objetos amontonados?", "Al reaparecer, cada objeto se coloca detrás de la menor coordenada Y con separación fija.")
qa("¿Qué es Room?", "Una capa oficial sobre SQLite local con validación y generación de código.")
qa("¿Usa PostgreSQL?", "No. CatchFood usa SQLite dentro del teléfono.")
qa("¿Qué hace el DAO?", "Define inserciones y consultas SQL de la tabla de puntajes.")
qa("¿Qué hace Repository?", "Oculta Room al ViewModel y expone operaciones del dominio.")
qa("¿Por qué suspend?", "Porque el acceso a almacenamiento puede esperar y no debe bloquear la UI.")
qa("¿Qué hace KSP?", "Genera en compilación la implementación necesaria de Room.")
qa("¿Por qué Singleton para Database?", "Para reutilizar una única conexión segura en el proceso.")
qa("¿Qué hace @Volatile?", "Hace visible entre hilos el valor más reciente de la instancia.")
qa("¿Qué hace synchronized?", "Evita que dos hilos creen la base simultáneamente.")
qa("¿Por qué scoreSavedForCurrentGame?", "Evita insertar la misma partida más de una vez.")
qa("¿Qué diferencia hay entre remember y ViewModel?", "remember conserva durante recomposición; ViewModel está diseñado para estado de pantalla y cambios de configuración.")
qa("¿Qué hace DisposableEffect?", "Conecta creación y limpieza de un recurso al ciclo de composición.")
qa("¿Qué mejora harías?", "Usaría enum para playerState y Flow para observar el ranking automáticamente.", "También añadiría migraciones de Room y pruebas unitarias para colisiones/spawn.")
doc.add_page_break()

# Mock defense
doc.add_heading("21. Simulacro oral", level=1)
doc.add_heading("Presentación de 3 minutos", level=2)
para("Mi módulo se llama CatchFood. Está construido con Jetpack Compose y tiene tres pantallas internas: menú, juego e histórico. La navegación está centralizada en CatchFoodNavHost y las pantallas reciben callbacks para no depender directamente del NavController.")
para("Durante el juego, CatchFoodViewModel coordina el estado visible y una clase CatchFoodGameLogic que contiene las reglas. CatchFoodUiState es una data class inmutable con score, golpes de veneno, comida perdida, estado visual, pausa y game over. La UI observa este objeto y se recompone únicamente cuando cambian valores visibles.")
para("El loop usa withFrameNanos para sincronizarse con Android. Calculo delta time para obtener movimiento consistente entre dispositivos. Los objetos se dibujan en un Canvas para reducir recomposición y layout. Las colisiones son AABB con hitboxes reducidas para ajustarse al contenido visible de los PNG.")
para("Cuando termina la partida, el ViewModel captura el puntaje una sola vez y lo guarda con una corrutina. El Repository utiliza Room sobre SQLite. La Entity representa la tabla, el DAO define inserción y consulta, y el histórico ejecuta un ORDER BY descendente con LIMIT 10. Finalmente, produceState convierte la consulta suspendida en estado observable y LazyColumn muestra el ranking.")

doc.add_heading("Si te interrumpen", level=2)
bullet("Si preguntan UI: hablá de Compose, State y recomposición.")
bullet("Si preguntan rendimiento: hablá de Canvas, frameVersion, sampling y delta time.")
bullet("Si preguntan arquitectura: hablá de callbacks, ViewModel, UiState y Repository.")
bullet("Si preguntan base: Entity, DAO, Database, Repository y SQLite.")
bullet("Si preguntan seguridad de recursos: remember, DisposableEffect y release.")

doc.add_heading("Errores que no debés decir", level=2)
bullet("No digás que Room es PostgreSQL.")
bullet("No digás que una corrutina es exactamente un hilo.")
bullet("No digás que remember guarda datos después de cerrar la app.")
bullet("No digás que Canvas evita toda recomposición; reduce el trabajo animado y separa redraw de layout.")
bullet("No digás que ViewModel dibuja la UI; administra estado y acciones.")
doc.add_page_break()

# Cheat sheet
doc.add_heading("22. Hoja de repaso final", level=1)
callout("Lambda:", "función anónima pasada como dato: `{ onGoToGame() }`.", LIGHT_PINK)
callout("Callback:", "lambda que otro componente llama cuando ocurre un evento.", LIGHT_PINK)
callout("UiState:", "fotografía inmutable de valores visibles.", LIGHT_PINK)
callout("ViewModel:", "coordina estado, eventos, motor y persistencia.", LIGHT_PINK)
callout("Composable:", "función declarativa que describe interfaz.", LIGHT_BLUE)
callout("Recomposición:", "actualización selectiva por cambios en State.", LIGHT_BLUE)
callout("remember:", "conserva un cálculo/objeto entre recomposiciones.", LIGHT_BLUE)
callout("Corrutina:", "trabajo suspendible que evita bloquear UI.", LIGHT_GREEN)
callout("Room:", "capa sobre SQLite local.", LIGHT_GREEN)
callout("DAO:", "contrato de acceso y consultas a datos.", LIGHT_GREEN)
callout("Canvas:", "dibujo directo eficiente para la escena animada.", LIGHT_YELLOW)
callout("Delta time:", "tiempo real entre frames.", LIGHT_YELLOW)
callout("AABB:", "superposición de rectángulos para colisiones.", LIGHT_YELLOW)

doc.add_heading("Diez frases que tenés que memorizar", level=2)
reset_numbering()
for sentence in [
    "La UI envía eventos y observa estado; no ejecuta SQL.",
    "El ViewModel coordina el juego y sobrevive a recomposiciones.",
    "UiState agrupa datos visibles en una data class inmutable.",
    "GameLogic contiene reglas, posiciones, colisiones y spawn.",
    "withFrameNanos sincroniza el loop con Android.",
    "Delta time mantiene velocidad consistente entre dispositivos.",
    "Canvas reduce el trabajo de animar múltiples elementos Compose.",
    "Las hitboxes reducidas evitan capturas visualmente falsas.",
    "Room usa SQLite local y KSP genera implementaciones.",
    "El Repository desacopla ViewModel de DAO y Database.",
]:
    numbered(sentence)

callout("Cierre recomendado:", "El diseño separa responsabilidades, optimiza el renderizado en Android real y conserva el historial local sin bloquear el hilo principal.", LIGHT_GREEN)

doc.core_properties.title = "Guía de defensa técnica - CatchFood"
doc.core_properties.subject = "Kotlin, Jetpack Compose, ViewModel, Canvas y Room"
doc.core_properties.author = "PlayLand"
doc.core_properties.keywords = "CatchFood, Kotlin, Compose, ViewModel, Room, SQLite, defensa"
normalize_page_breaks_and_numbering(doc)
doc.save(DOCX_PATH)
print(DOCX_PATH)
