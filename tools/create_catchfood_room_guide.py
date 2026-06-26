from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\geral\Downloads\PlayLand\Playland")
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)
DOCX_PATH = OUT / "Guia_Room_CatchFood.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
PINK = RGBColor(0xE9, 0x6E, 0x9D)
GREEN = RGBColor(0x4F, 0x8D, 0x3A)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT_BLUE = "E8EEF5"
LIGHT_PINK = "FCE8F0"
LIGHT_GREEN = "EAF4E5"
LIGHT_YELLOW = "FFF4CC"
CODE_FILL = "F4F6F9"


def set_font(run, name="Calibri", size=11, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def set_table_geometry(table, widths_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_inches) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            dxa = int(widths_inches[idx] * 1440)
            cell.width = Inches(widths_inches[idx])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(dxa))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "PLAYLAND · CATCHFOOD · GUÍA ROOM / SQLITE"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.runs[0], size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


doc = Document()
style_doc(doc)


def title(text, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(82)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=28, bold=True, color=PINK)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(22)
    r2 = p2.add_run(subtitle)
    set_font(r2, size=14, color=DARK_BLUE)


def para(text="", bold_start=None, italic=False, color=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_start and text.startswith(bold_start):
        r1 = p.add_run(bold_start)
        set_font(r1, bold=True, color=color or DARK)
        r2 = p.add_run(text[len(bold_start):])
        set_font(r2, italic=italic, color=color or DARK)
    else:
        r = p.add_run(text)
        set_font(r, italic=italic, color=color or DARK)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r)
    return p


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_font(r)
    return p


def code_block(text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.35])
    cell = table.cell(0, 0)
    shade(cell, CODE_FILL)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(text.strip("\n").splitlines()):
        if idx:
            p.add_run("\n")
        r = p.add_run(line)
        set_font(r, name="Consolas", size=8.5, color=DARK)
    para("", after=2)


def callout(label, text, fill=LIGHT_YELLOW):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.35])
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r1 = p.add_run(label + ": ")
    set_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_font(r2, color=DARK)
    para("", after=2)


def simple_table(headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        shade(cell, header_fill)
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_font(r, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            shade(cells[i], "FFFFFF")
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(value)
            set_font(r, size=10.2, color=DARK)
    para("", after=4)
    return table


title(
    "Guía Room en CatchFood",
    "Cómo se guardó el histórico TOP 10 usando Room sobre SQLite"
)

callout(
    "Idea central para decir en defensa",
    "En CatchFood usé Room para guardar localmente el puntaje final de cada partida. Room genera el código que conecta Kotlin con SQLite, y luego la pantalla de histórico consulta los 10 mejores valores ordenados de mayor a menor.",
    LIGHT_PINK,
)

para("Esta guía se enfoca únicamente en la persistencia de datos de CatchFood: qué problema resuelve, qué archivos se agregaron, cómo viaja el puntaje desde el juego hasta SQLite y cómo responder preguntas típicas del jurado.")

doc.add_heading("1. Primero: Room, SQLite y PostgreSQL no son lo mismo", level=1)
simple_table(
    ["Concepto", "Qué es", "Cómo aplica aquí"],
    [
        ("SQLite", "Base de datos local que vive dentro del teléfono.", "Aquí guarda el archivo catch_food.db con los puntajes."),
        ("Room", "Librería de Android que simplifica el uso de SQLite.", "Permite usar @Entity, @Dao y @Database en vez de escribir todo SQL manual."),
        ("PostgreSQL", "Motor de base de datos de servidor.", "No se usa en CatchFood. Sería para backend o nube, no para este histórico local."),
    ],
    [1.25, 2.55, 2.55],
)
callout(
    "Respuesta corta",
    "Room no es PostgreSQL. Room trabaja sobre SQLite local en Android. Lo escogí porque el histórico del juego no necesita internet ni servidor.",
    LIGHT_GREEN,
)

doc.add_heading("2. Archivos donde se aplicó Room", level=1)
simple_table(
    ["Archivo", "Responsabilidad"],
    [
        ("CatchFoodScoreEntity.kt", "Define la tabla catch_food_scores y sus columnas."),
        ("CatchFoodScoreDao.kt", "Define las operaciones SQL: insertar puntaje y consultar el TOP 10."),
        ("CatchFoodDatabase.kt", "Crea la base de datos Room y expone el DAO."),
        ("CatchFoodScoreRepository.kt", "Capa intermedia para que la UI/ViewModel no hable directo con Room."),
        ("CatchFoodViewModel.kt", "Guarda el puntaje cuando el juego termina."),
        ("CatchFoodLeaderboard.kt", "Lee los puntajes guardados y los muestra en pantalla."),
        ("app/build.gradle.kts", "Declara dependencias Room y KSP."),
    ],
    [2.2, 4.15],
)

doc.add_heading("3. Configuración Gradle: dependencias necesarias", level=1)
para("En app/build.gradle.kts se agregaron las dependencias de Room y el procesador KSP. KSP genera código en compilación para que Room pueda crear la implementación real de la base de datos y del DAO.")
code_block(
    """
plugins {
    alias(libs.plugins.ksp)
}

dependencies {
    implementation(libs.androidx.room.runtime)
    implementation(libs.androidx.room.ktx)
    ksp(libs.androidx.room.compiler)
}
"""
)
bullet("room-runtime: contiene las clases principales de Room.")
bullet("room-ktx: agrega soporte cómodo para Kotlin y corrutinas suspend.")
bullet("room-compiler con ksp: analiza las anotaciones @Entity, @Dao y @Database, y genera código automáticamente.")

doc.add_heading("4. Entity: la tabla de puntajes", level=1)
para("La Entity es la clase que Room convierte en una tabla SQLite. En CatchFood se llama CatchFoodScoreEntity.")
code_block(
    """
@Entity(tableName = "catch_food_scores")
data class CatchFoodScoreEntity(
    @PrimaryKey(autoGenerate = true)
    val id: Long = 0,
    val foodCaught: Int
)
"""
)
simple_table(
    ["Parte", "Explicación para defensa"],
    [
        ("@Entity", "Le dice a Room: esta clase representa una tabla."),
        ("tableName", "Define el nombre real de la tabla en SQLite: catch_food_scores."),
        ("data class", "Clase de Kotlin pensada para guardar datos. Genera equals, copy, toString, etc."),
        ("@PrimaryKey", "Marca la columna id como identificador único de cada fila."),
        ("autoGenerate = true", "SQLite/Room asigna el id automáticamente al insertar."),
        ("foodCaught", "Es el dato importante: la cantidad de comida que atrapó el jugador."),
    ],
    [1.8, 4.55],
)
callout("Cómo se ve como tabla", "id sería una columna numérica única; foodCaught sería otra columna numérica con el puntaje final.", LIGHT_BLUE)

doc.add_heading("5. DAO: las consultas de la base", level=1)
para("DAO significa Data Access Object. Es una interfaz donde se declaran las operaciones permitidas contra la base. Room genera la implementación real.")
code_block(
    """
@Dao
interface CatchFoodScoreDao {
    @Insert
    suspend fun insert(score: CatchFoodScoreEntity)

    @Query(
        \"\"\"
        SELECT foodCaught
        FROM catch_food_scores
        ORDER BY foodCaught DESC, id ASC
        LIMIT 10
        \"\"\"
    )
    suspend fun getTopTenScores(): List<Int>
}
"""
)
bullet("@Dao marca la interfaz como una puerta de acceso a la base.")
bullet("@Insert indica que la función insertará una fila en la tabla.")
bullet("suspend permite ejecutar la operación dentro de una corrutina, sin bloquear la pantalla del juego.")
bullet("@Query contiene SQL puro para consultar la tabla.")
bullet("ORDER BY foodCaught DESC ordena de mayor puntaje a menor puntaje.")
bullet("id ASC desempata: si dos partidas tienen el mismo puntaje, aparece primero la que se guardó antes.")
bullet("LIMIT 10 hace que solo salgan los 10 mejores resultados.")

doc.add_heading("6. Database: creación de Room y patrón Singleton", level=1)
code_block(
    """
@Database(
    entities = [CatchFoodScoreEntity::class],
    version = 1,
    exportSchema = false
)
abstract class CatchFoodDatabase : RoomDatabase() {
    abstract fun scoreDao(): CatchFoodScoreDao
}
"""
)
para("La clase CatchFoodDatabase representa la base de datos completa de CatchFood. Hereda de RoomDatabase porque Room necesita una clase central donde se registren las entidades y DAOs.")
simple_table(
    ["Elemento", "Qué significa"],
    [
        ("entities", "Lista de tablas manejadas por esta base. Aquí solo hay una: CatchFoodScoreEntity."),
        ("version = 1", "Versión inicial del esquema. Si cambia la estructura de tablas en el futuro, se sube."),
        ("exportSchema = false", "Evita exportar archivos de esquema. En proyectos grandes puede ponerse true."),
        ("abstract fun scoreDao()", "Room implementa esta función para devolver el DAO real."),
    ],
    [1.75, 4.6],
)

code_block(
    """
companion object {
    @Volatile
    private var instance: CatchFoodDatabase? = null

    fun getInstance(context: Context): CatchFoodDatabase = instance ?: synchronized(this) {
        instance ?: Room.databaseBuilder(
            context.applicationContext,
            CatchFoodDatabase::class.java,
            DATABASE_NAME
        ).build().also { instance = it }
    }
}
"""
)
bullet("@Volatile ayuda a que el valor instance sea visible correctamente entre hilos.")
bullet("instance guarda una única copia de la base en memoria.")
bullet("synchronized evita que dos hilos creen dos bases al mismo tiempo.")
bullet("context.applicationContext evita guardar accidentalmente una Activity y causar fugas de memoria.")
bullet("Room.databaseBuilder construye la base SQLite local.")
bullet("DATABASE_NAME = \"catch_food.db\" es el nombre del archivo físico de la base.")

doc.add_heading("7. Repository: capa intermedia limpia", level=1)
code_block(
    """
class CatchFoodScoreRepository(context: Context) {
    private val scoreDao = CatchFoodDatabase.getInstance(context).scoreDao()

    suspend fun saveScore(score: Int) {
        scoreDao.insert(CatchFoodScoreEntity(foodCaught = score))
    }

    suspend fun getTopScores(): List<Int> = scoreDao.getTopTenScores()
}
"""
)
para("El Repository existe para que el resto del juego no tenga que conocer detalles de Room. Si mañana se cambia la base local por una API, el cambio se concentra aquí.")
bullet("saveScore recibe un Int normal del juego y lo convierte en CatchFoodScoreEntity.")
bullet("getTopScores devuelve una lista de Int porque la pantalla solo necesita mostrar el valor foodCaught.")
bullet("Ambas funciones son suspend porque terminan llamando operaciones de base de datos.")

doc.add_heading("8. ViewModel: guardar el puntaje al perder", level=1)
code_block(
    """
private val scoreRepository = CatchFoodScoreRepository(application)
private var scoreSavedForCurrentGame = false

fun updateGame(deltaSeconds: Float) {
    gameLogic.updateGame(deltaSeconds)
    syncUiState()

    if (gameLogic.isGameOver && !scoreSavedForCurrentGame) {
        scoreSavedForCurrentGame = true
        val finalScore = gameLogic.score
        viewModelScope.launch {
            scoreRepository.saveScore(finalScore)
        }
    }
}
"""
)
bullet("El ViewModel detecta cuando gameLogic.isGameOver cambia a verdadero.")
bullet("scoreSavedForCurrentGame evita guardar el mismo resultado muchas veces. Esto es clave porque updateGame se ejecuta repetidamente durante el loop.")
bullet("finalScore copia el puntaje final antes de entrar a la corrutina.")
bullet("viewModelScope.launch abre una corrutina asociada al ciclo de vida del ViewModel.")
bullet("scoreRepository.saveScore(finalScore) inserta el registro en SQLite usando Room.")
callout("Frase para defensa", "Guardé el puntaje desde el ViewModel, no desde el Composable directamente, para separar la lógica del juego y la persistencia de la interfaz visual.", LIGHT_GREEN)

doc.add_heading("9. Pantalla histórico: leer el TOP 10", level=1)
code_block(
    """
val context = LocalContext.current
val scoreRepository = remember(context) { CatchFoodScoreRepository(context) }
val scores by produceState(initialValue = emptyList(), scoreRepository) {
    value = scoreRepository.getTopScores()
}
"""
)
bullet("LocalContext.current da acceso al Context de Android desde Compose.")
bullet("remember evita recrear el Repository en cada recomposición innecesaria.")
bullet("produceState crea un estado de Compose a partir de una operación suspend.")
bullet("initialValue = emptyList() muestra una lista vacía mientras se carga la consulta.")
bullet("value = scoreRepository.getTopScores() consulta Room y actualiza la UI automáticamente.")

code_block(
    """
LazyColumn {
    itemsIndexed(scores) { index, score ->
        Text(text = \"${index + 1}.\")
        Text(text = score.toString())
    }
}
"""
)
para("LazyColumn muestra la lista de puntajes de forma eficiente. itemsIndexed entrega tanto la posición como el valor, por eso se puede enseñar 1, 2, 3... junto al puntaje.")

doc.add_heading("10. Flujo completo del dato", level=1)
simple_table(
    ["Paso", "Qué ocurre", "Archivo"],
    [
        ("1", "El jugador juega y consigue comida.", "CatchFoodGameLogic.kt"),
        ("2", "El juego termina por veneno o pérdidas.", "CatchFoodViewModel.kt"),
        ("3", "El ViewModel detecta isGameOver.", "CatchFoodViewModel.kt"),
        ("4", "Se toma gameLogic.score como puntaje final.", "CatchFoodViewModel.kt"),
        ("5", "Se llama saveScore en una corrutina.", "CatchFoodViewModel.kt"),
        ("6", "El Repository crea CatchFoodScoreEntity.", "CatchFoodScoreRepository.kt"),
        ("7", "El DAO inserta la fila en catch_food_scores.", "CatchFoodScoreDao.kt"),
        ("8", "SQLite guarda el dato en catch_food.db.", "CatchFoodDatabase.kt"),
        ("9", "Histórico consulta SELECT foodCaught ... LIMIT 10.", "CatchFoodScoreDao.kt"),
        ("10", "Compose muestra el TOP 10 en pantalla.", "CatchFoodLeaderboard.kt"),
    ],
    [0.55, 3.45, 2.35],
    LIGHT_GREEN,
)

doc.add_heading("11. Por qué se usó así y no de otra forma", level=1)
bullet("Se usa Room porque el histórico es local, pequeño y no necesita servidor.")
bullet("Se usa SQLite porque ya viene integrado en Android.")
bullet("Se usa Repository para separar acceso a datos de la pantalla.")
bullet("Se usa ViewModel para que la lógica sobreviva mejor a recomposiciones y no dependa directamente del Composable.")
bullet("Se usan corrutinas porque las operaciones de base no deben bloquear el hilo principal.")
bullet("Se consulta solo List<Int> porque el requerimiento era mostrar únicamente la comida máxima alcanzada, no fecha ni nombre.")

doc.add_heading("12. Preguntas típicas de defensa", level=1)
simple_table(
    ["Pregunta", "Respuesta recomendada"],
    [
        ("¿Qué es Room?", "Una librería de Android que facilita trabajar con SQLite usando clases, anotaciones y DAOs."),
        ("¿Room es PostgreSQL?", "No. Room trabaja sobre SQLite local. PostgreSQL es una base de datos de servidor."),
        ("¿Dónde se crea la tabla?", "En CatchFoodScoreEntity con @Entity(tableName = \"catch_food_scores\")."),
        ("¿Dónde está el SQL?", "En CatchFoodScoreDao, dentro de la anotación @Query."),
        ("¿Por qué suspend?", "Porque insertar o consultar base de datos debe hacerse fuera del hilo principal usando corrutinas."),
        ("¿Por qué hay Repository?", "Para separar la lógica de datos de la UI y del ViewModel."),
        ("¿Qué evita scoreSavedForCurrentGame?", "Evita guardar el mismo puntaje varias veces mientras el estado Game Over sigue activo."),
        ("¿Qué pasa si cierro la app?", "El dato queda guardado en SQLite porque es persistencia local."),
    ],
    [2.1, 4.25],
    LIGHT_PINK,
)

doc.add_heading("13. Mini guion para explicarlo en 60 segundos", level=1)
para("“Para el histórico de CatchFood implementé persistencia local usando Room, que es una capa de Android encima de SQLite. Definí una entidad llamada CatchFoodScoreEntity, que representa la tabla catch_food_scores con dos columnas: id y foodCaught. Luego creé un DAO con dos operaciones: insertar un puntaje y consultar los 10 mejores con SQL ordenado por foodCaught descendente. La clase CatchFoodDatabase construye la base catch_food.db con Room.databaseBuilder y usa Singleton para no crear múltiples instancias. Encima puse un Repository para que el ViewModel no dependa directamente de Room. Cuando el juego termina, el ViewModel detecta isGameOver, toma el score final y lo guarda una sola vez con viewModelScope.launch. Finalmente, la pantalla CatchFoodLeaderboard usa produceState para pedirle al Repository el TOP 10 y mostrarlo en Compose.”")

doc.add_heading("14. Checklist mental antes de defender", level=1)
bullet("Room = librería; SQLite = base local; PostgreSQL = servidor externo no usado.")
bullet("@Entity define tabla.")
bullet("@PrimaryKey(autoGenerate = true) crea id automático.")
bullet("@Dao define operaciones.")
bullet("@Insert inserta.")
bullet("@Query ejecuta SQL.")
bullet("@Database registra entidades y versión.")
bullet("Repository encapsula acceso a datos.")
bullet("ViewModel guarda cuando pierde.")
bullet("Leaderboard consulta y muestra el TOP 10.")

doc.add_heading("15. Rutas exactas para enseñar en Android Studio", level=1)
for path in [
    "app/src/main/java/com/example/playland2/feature/catchfood/data/CatchFoodScoreEntity.kt",
    "app/src/main/java/com/example/playland2/feature/catchfood/data/CatchFoodScoreDao.kt",
    "app/src/main/java/com/example/playland2/feature/catchfood/data/CatchFoodDatabase.kt",
    "app/src/main/java/com/example/playland2/feature/catchfood/data/CatchFoodScoreRepository.kt",
    "app/src/main/java/com/example/playland2/feature/catchfood/ui/screen/CatchFoodViewModel.kt",
    "app/src/main/java/com/example/playland2/feature/catchfood/ui/screen/CatchFoodLeaderboard.kt",
    "app/build.gradle.kts",
]:
    code_block(path)

doc.core_properties.title = "Guía Room en CatchFood"
doc.core_properties.subject = "Persistencia local con Room y SQLite en CatchFood"
doc.core_properties.author = "Codex"
doc.save(DOCX_PATH)
print(DOCX_PATH)
